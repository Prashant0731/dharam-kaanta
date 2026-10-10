from fastapi import FastAPI, HTTPException, Depends, Query, Response
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from typing import Optional, List
from datetime import datetime, date
import sqlite3
import uuid
import json
from contextlib import asynccontextmanager
from io import BytesIO

# Database setup
DB_PATH = "dharam_kaanta.db"

def init_db():
    conn = sqlite3.connect(DB_PATH)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS entries (
            id TEXT PRIMARY KEY,
            sno INTEGER NOT NULL,
            vehicle TEXT,
            tare TEXT,
            gross TEXT,
            net REAL DEFAULT 0,
            money REAL DEFAULT 0,
            vehicle_photo TEXT,
            gross_photo TEXT,
            saved INTEGER DEFAULT 0,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP,
            updated_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    """)
    conn.commit()
    conn.close()

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    try:
        yield conn
    finally:
        conn.close()

@asynccontextmanager
async def lifespan(app: FastAPI):
    init_db()
    print("Starting Dharam Kaanta API with SQLite...")
    yield
    print("Shutting down...")

app = FastAPI(title="Dharam Kaanta API", version="1.0.0", lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Pydantic models
class EntryBase(BaseModel):
    sno: int
    vehicle: Optional[str] = None
    tare: Optional[str] = None
    gross: Optional[str] = None
    vehiclePhoto: Optional[str] = None
    grossPhoto: Optional[str] = None

class EntryCreate(EntryBase):
    pass

class EntryUpdate(BaseModel):
    vehicle: Optional[str] = None
    tare: Optional[str] = None
    gross: Optional[str] = None
    vehiclePhoto: Optional[str] = None
    grossPhoto: Optional[str] = None
    saved: Optional[bool] = None

class EntryResponse(EntryBase):
    id: str
    net: float
    money: float
    saved: bool
    createdAt: Optional[str] = None
    updatedAt: Optional[str] = None

class StatsResponse(BaseModel):
    totalEntries: int
    savedEntries: int
    todayEntries: int
    totalWeight: float
    totalRevenue: float

# Helpers
def calc_net_money(tare: str, gross: str):
    t = float(tare) if tare else 0
    g = float(gross) if gross else 0
    net = g - t if (t or g) else 0
    money = round(net / 5, 2) if net else 0
    return net, money

# Routes
@app.get("/")
def root():
    return {"name": "Dharam Kaanta API", "version": "1.0.0", "docs": "/docs"}

@app.get("/health")
def health():
    return {"status": "ok"}

@app.get("/api/v1/entries", response_model=List[EntryResponse])
def get_entries(
    skip: int = 0, limit: int = 100, 
    search: Optional[str] = None,
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    db: sqlite3.Connection = Depends(get_db)
):
    query = "SELECT * FROM entries WHERE 1=1"
    params = []
    
    if search:
        query += " AND vehicle LIKE ?"
        params.append(f"%{search}%")
    if date_from:
        query += " AND created_at >= ?"
        params.append(date_from)
    if date_to:
        query += " AND created_at <= ?"
        params.append(date_to)
    
    query += " ORDER BY created_at DESC LIMIT ? OFFSET ?"
    params.extend([limit, skip])
    
    rows = db.execute(query, params).fetchall()
    return [dict(row) for row in rows]

@app.get("/api/v1/entries/{entry_id}", response_model=EntryResponse)
def get_entry(entry_id: str, db: sqlite3.Connection = Depends(get_db)):
    row = db.execute("SELECT * FROM entries WHERE id = ?", (entry_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Entry not found")
    return dict(row)

@app.post("/api/v1/entries", response_model=EntryResponse)
def create_entry(entry: EntryCreate, db: sqlite3.Connection = Depends(get_db)):
    net, money = calc_net_money(entry.tare, entry.gross)
    entry_id = str(uuid.uuid4())
    now = datetime.now().isoformat()
    
    db.execute("""
        INSERT INTO entries (id, sno, vehicle, tare, gross, net, money, vehicle_photo, gross_photo, saved, created_at, updated_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
    """, (entry_id, entry.sno, entry.vehicle, entry.tare, entry.gross, net, money, 
          entry.vehiclePhoto, entry.grossPhoto, 0, now, now))
    db.commit()
    
    row = db.execute("SELECT * FROM entries WHERE id = ?", (entry_id,)).fetchone()
    return dict(row)

@app.put("/api/v1/entries/{entry_id}", response_model=EntryResponse)
def update_entry(entry_id: str, entry: EntryUpdate, db: sqlite3.Connection = Depends(get_db)):
    existing = db.execute("SELECT * FROM entries WHERE id = ?", (entry_id,)).fetchone()
    if not existing:
        raise HTTPException(404, "Entry not found")
    
    update_data = entry.model_dump(exclude_unset=True)
    net, money = calc_net_money(
        update_data.get('tare', existing['tare']),
        update_data.get('gross', existing['gross'])
    )
    update_data['net'] = net
    update_data['money'] = money
    update_data['updated_at'] = datetime.now().isoformat()
    
    if 'saved' in update_data:
        update_data['saved'] = 1 if update_data['saved'] else 0
    
    fields = ", ".join([f"{k} = ?" for k in update_data.keys()])
    values = list(update_data.values()) + [entry_id]
    db.execute(f"UPDATE entries SET {fields} WHERE id = ?", values)
    db.commit()
    
    row = db.execute("SELECT * FROM entries WHERE id = ?", (entry_id,)).fetchone()
    return dict(row)

@app.delete("/api/v1/entries/{entry_id}")
def delete_entry(entry_id: str, db: sqlite3.Connection = Depends(get_db)):
    db.execute("DELETE FROM entries WHERE id = ?", (entry_id,))
    db.commit()
    return {"message": "Entry deleted"}

@app.post("/api/v1/entries/bulk", response_model=List[EntryResponse])
def bulk_create(entries: List[EntryCreate], db: sqlite3.Connection = Depends(get_db)):
    created = []
    for entry in entries:
        net, money = calc_net_money(entry.tare, entry.gross)
        entry_id = str(uuid.uuid4())
        now = datetime.now().isoformat()
        
        db.execute("""
            INSERT INTO entries (id, sno, vehicle, tare, gross, net, money, vehicle_photo, gross_photo, saved, created_at, updated_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (entry_id, entry.sno, entry.vehicle, entry.tare, entry.gross, net, money,
              entry.vehiclePhoto, entry.grossPhoto, 0, now, now))
        created.append(entry_id)
    db.commit()
    
    placeholders = ",".join("?" * len(created))
    rows = db.execute(f"SELECT * FROM entries WHERE id IN ({placeholders})", created).fetchall()
    return [dict(row) for row in rows]

@app.get("/api/v1/stats", response_model=StatsResponse)
def get_stats(db: sqlite3.Connection = Depends(get_db)):
    total = db.execute("SELECT COUNT(*) FROM entries").fetchone()[0]
    saved = db.execute("SELECT COUNT(*) FROM entries WHERE saved = 1").fetchone()[0]
    today = db.execute("SELECT COUNT(*) FROM entries WHERE date(created_at) = date('now')").fetchone()[0]
    total_weight = db.execute("SELECT COALESCE(SUM(net), 0) FROM entries").fetchone()[0]
    total_revenue = db.execute("SELECT COALESCE(SUM(money), 0) FROM entries").fetchone()[0]
    
    return StatsResponse(
        totalEntries=total, savedEntries=saved, todayEntries=today,
        totalWeight=total_weight, totalRevenue=total_revenue
    )

# Export endpoints
@app.get("/api/v1/export")
def export_data(
    format: str = Query(..., pattern="^(excel|pdf|word)$"),
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    vehicle_filter: Optional[str] = None,
    db: sqlite3.Connection = Depends(get_db)
):
    query = "SELECT * FROM entries WHERE 1=1"
    params = []
    
    if vehicle_filter:
        query += " AND vehicle LIKE ?"
        params.append(f"%{vehicle_filter}%")
    if date_from:
        query += " AND created_at >= ?"
        params.append(date_from)
    if date_to:
        query += " AND created_at <= ?"
        params.append(date_to)
    
    query += " ORDER BY created_at DESC"
    entries = db.execute(query, params).fetchall()
    
    if not entries:
        raise HTTPException(404, "No data to export")
    
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"dharam_kaanta_{timestamp}"
    
    if format == "excel":
        return export_excel(entries, filename)
    elif format == "pdf":
        return export_pdf(entries, filename)
    elif format == "word":
        return export_word(entries, filename)

def export_excel(entries, filename):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Entries"
    
    headers = ['S.No', 'Vehicle No.', 'Tare (kg)', 'Gross (kg)', 'Net (kg)', 
               'Rupees', 'Vehicle Photo', 'Gross Photo', 'Status', 'Created At']
    
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="17202A", end_color="17202A", fill_type="solid")
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    
    for row_idx, entry in enumerate(entries, 2):
        ws.cell(row=row_idx, column=1, value=row_idx - 1)
        ws.cell(row=row_idx, column=2, value=entry['vehicle'] or "")
        ws.cell(row=row_idx, column=3, value=entry['tare'] or "")
        ws.cell(row=row_idx, column=4, value=entry['gross'] or "")
        ws.cell(row=row_idx, column=5, value=entry['net'] if entry['net'] else "")
        ws.cell(row=row_idx, column=6, value=round(entry['money'], 2) if entry['money'] else "")
        ws.cell(row=row_idx, column=7, value="Yes" if entry['vehicle_photo'] else "No")
        ws.cell(row=row_idx, column=8, value="Yes" if entry['gross_photo'] else "No")
        ws.cell(row=row_idx, column=9, value="Saved" if entry['saved'] else "Draft")
        ws.cell(row=row_idx, column=10, value=entry['created_at'] or "")
    
    for i, width in enumerate([8, 20, 14, 14, 14, 14, 16, 16, 12, 22], 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
    
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}.xlsx"'}
    )

def export_pdf(entries, filename):
    from fpdf import FPDF
    
    class PDF(FPDF):
        def header(self):
            self.set_font('Helvetica', 'B', 16)
            self.cell(0, 10, 'Dharam Kaanta - Weighbridge Entries', 0, 1, 'C')
            self.set_font('Helvetica', '', 8)
            self.cell(0, 5, f'Generated: {datetime.now().strftime("%d-%m-%Y %H:%M")}', 0, 1, 'C')
            self.cell(0, 5, f'Total Entries: {len(entries)}', 0, 1, 'C')
            self.ln(5)
    
    pdf = PDF('L', 'mm', 'A4')
    pdf.add_page()
    pdf.set_font('Helvetica', 'B', 7)
    
    headers = ['S.No', 'Vehicle', 'Tare', 'Gross', 'Net', 'Rupees', 'V-Photo', 'G-Photo', 'Status', 'Created']
    col_widths = [10, 35, 18, 18, 18, 18, 16, 16, 16, 30]
    
    pdf.set_fill_color(23, 32, 42)
    pdf.set_text_color(255, 255, 255)
    for i, header in enumerate(headers):
        pdf.cell(col_widths[i], 7, header, 1, 0, 'C', True)
    pdf.ln()
    
    pdf.set_text_color(0, 0, 0)
    pdf.set_font('Helvetica', '', 7)
    fill = False
    for idx, entry in enumerate(entries):
        pdf.set_fill_color(243, 245, 247) if fill else pdf.set_fill_color(255, 255, 255)
        
        pdf.cell(col_widths[0], 6, str(idx + 1), 1, 0, 'C', True)
        pdf.cell(col_widths[1], 6, entry['vehicle'] or "-", 1, 0, 'L', True)
        pdf.cell(col_widths[2], 6, entry['tare'] or "-", 1, 0, 'C', True)
        pdf.cell(col_widths[3], 6, entry['gross'] or "-", 1, 0, 'C', True)
        pdf.cell(col_widths[4], 6, str(entry['net']) if entry['net'] else "-", 1, 0, 'C', True)
        pdf.cell(col_widths[5], 6, f"{entry['money']:.2f}" if entry['money'] else "-", 1, 0, 'C', True)
        pdf.cell(col_widths[6], 6, "Yes" if entry['vehicle_photo'] else "No", 1, 0, 'C', True)
        pdf.cell(col_widths[7], 6, "Yes" if entry['gross_photo'] else "No", 1, 0, 'C', True)
        pdf.cell(col_widths[8], 6, "Saved" if entry['saved'] else "Draft", 1, 0, 'C', True)
        pdf.cell(col_widths[9], 6, entry['created_at'] or "-", 1, 0, 'C', True)
        pdf.ln()
        fill = not fill
    
    buffer = BytesIO()
    pdf.output(buffer)
    buffer.seek(0)
    
    return Response(
        content=buffer.getvalue(),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}.pdf"'}
    )

def export_word(entries, filename):
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    
    doc = Document()
    
    for section in doc.sections:
        section.orientation = 1
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)
    
    title = doc.add_heading('Dharam Kaanta - Weighbridge Entries', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(f'Generated: {datetime.now().strftime("%d-%m-%Y %H:%M")} | Total Entries: {len(entries)}')
    run.font.size = Pt(10)
    
    table = doc.add_table(rows=1, cols=10)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    headers = ['S.No', 'Vehicle No.', 'Tare (kg)', 'Gross (kg)', 'Net (kg)', 
               'Rupees', 'Vehicle Photo', 'Gross Photo', 'Status', 'Created At']
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
        for p in table.rows[0].cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(8)
    
    for idx, entry in enumerate(entries):
        row = table.add_row().cells
        row[0].text = str(idx + 1)
        row[1].text = entry['vehicle'] or "-"
        row[2].text = entry['tare'] or "-"
        row[3].text = entry['gross'] or "-"
        row[4].text = str(entry['net']) if entry['net'] else "-"
        row[5].text = f"{entry['money']:.2f}" if entry['money'] else "-"
        row[6].text = "Yes" if entry['vehicle_photo'] else "No"
        row[7].text = "Yes" if entry['gross_photo'] else "No"
        row[8].text = "Saved" if entry['saved'] else "Draft"
        row[9].text = entry['created_at'] or "-"
        for cell in row:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(8)
    
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}.docx"'}
    )

if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)